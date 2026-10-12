import pptx

prs = pptx.Presentation(r'C:\Users\Dharshini\Downloads\ppt.pptx')
print(f"Presentation Slide Width: {prs.slide_width/914400:.2f} in, Height: {prs.slide_height/914400:.2f} in")

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n=== SLIDE {idx} ===")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            tf = shape.text_frame
            full_text = " // ".join(p.text.strip() for p in tf.paragraphs if p.text.strip())
            print(f"  Shape {s_idx} ({shape.name}): L={shape.left/914400:.2f}, T={shape.top/914400:.2f}, W={shape.width/914400:.2f}, H={shape.height/914400:.2f}")
            print(f"     Content: {full_text[:120]}")
